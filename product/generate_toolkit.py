# -*- coding: utf-8 -*-
"""يولّد «عدة التحرر من الديون» — ملف Excel العربي القابل للبيع.

التشغيل:  python product/generate_toolkit.py
الناتج:   product/عدة-التحرر-من-الديون.xlsx
"""

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "product/عدة-التحرر-من-الديون.xlsx"

TEAL = "0F766E"
TEAL_DARK = "115E59"
MINT = "ECFDF5"
YELLOW = "FFF9C4"
GRAY = "F3F4F6"

F_TITLE = Font(name="Arial", size=16, bold=True, color="FFFFFF")
F_HEAD = Font(name="Arial", size=11, bold=True, color="FFFFFF")
F_SECTION = Font(name="Arial", size=12, bold=True, color=TEAL_DARK)
F_BASE = Font(name="Arial", size=11)
F_BOLD = Font(name="Arial", size=11, bold=True)
F_INPUT = Font(name="Arial", size=11, color="0000FF")          # مدخلات يعدّلها المستخدم
F_LINK = Font(name="Arial", size=11, color="008000")           # مرتبط بورقة أخرى
F_LINK_BOLD = Font(name="Arial", size=11, bold=True, color="008000")
F_NOTE = Font(name="Arial", size=9, italic=True, color="6B7280")
F_WARN = Font(name="Arial", size=11, bold=True, color="B91C1C")

FILL_TITLE = PatternFill("solid", fgColor=TEAL)
FILL_HEAD = PatternFill("solid", fgColor=TEAL_DARK)
FILL_SECTION = PatternFill("solid", fgColor=MINT)
FILL_INPUT = PatternFill("solid", fgColor=YELLOW)
FILL_ALT = PatternFill("solid", fgColor=GRAY)

THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CUR = '#,##0 "ر.س"'
PCT = "0.0%"
DATE = "yyyy/mm"

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

N_DEBTS = 10          # عدد خانات الديون في السجل
N_MONTHS = 120        # طول جدول السداد (10 سنوات)
SCHED_START = 12      # أول صف في جدول السداد
SCHED_END = SCHED_START + N_MONTHS - 1


def style(ws, cell, value=None, font=F_BASE, fill=None, fmt=None, align=RIGHT, border=True):
    c = ws[cell]
    if value is not None:
        c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align
    if border:
        c.border = BORDER
    return c


def title_row(ws, text, last_col="G", row=1):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = F_TITLE
    c.fill = FILL_TITLE
    c.alignment = CENTER
    ws.row_dimensions[row].height = 34


wb = Workbook()

# ============================================================ 1) ابدأ هنا
ws = wb.active
ws.title = "ابدأ هنا"
ws.sheet_view.rightToLeft = True
ws.sheet_properties.tabColor = TEAL
title_row(ws, "عدة التحرر من الديون", "E")
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 62
ws.column_dimensions["D"].width = 22
ws.column_dimensions["E"].width = 8

rows = [
    ("", ""),
    ("مرحباً بك!", "هذه العدة تحوّل ديونك من همّ غامض إلى خطة واضحة بتاريخ نهاية محدد. اتبع الخطوات الأربع بالترتيب."),
    ("", ""),
    ("الخطوة 1 — ورقة «الديون»", "سجّل كل دين: المبلغ المتبقي، الفائدة السنوية، والقسط الأدنى. الخلايا الصفراء هي أماكن الكتابة."),
    ("الخطوة 2 — ورقة «خطة السداد»", "حدّد الدفعة الشهرية الإضافية التي تستطيعها. سيُحسب تاريخ تحررك وإجمالي الفوائد تلقائياً."),
    ("الخطوة 3 — ورقة «الميزانية الشهرية»", "خطط دخلك ومصروفاتك الشهرية حتى تعرف من أين تأتي الدفعة الإضافية."),
    ("الخطوة 4 — ورقة «لوحة المتابعة»", "سجّل ما سددته فعلياً كل شهر وشاهد شريط الإنجاز يمتلئ."),
    ("", ""),
    ("دليل الألوان", ""),
    ("خلية صفراء بخط أزرق", "مدخلات — عدّلها بحرّية"),
    ("خط أسود", "معادلة تلقائية — لا تعدّلها"),
    ("خط أخضر", "قيمة مرتبطة بورقة أخرى — لا تعدّلها"),
    ("", ""),
    ("ملاحظة منهجية", "جدول السداد يستخدم متوسط الفائدة المرجّح لديونك مجتمعة (نموذج تقريبي موثّق في ورقة الخطة). "
     "للحصول على جدول دقيق لكل دين على حدة بطريقة كرة الثلج أو الانهيار الجليدي، استخدم الحاسبة المجانية المرافقة للعدة."),
]
r = 2
for label, desc in rows:
    if label:
        style(ws, f"B{r}", label, F_BOLD, border=False)
        style(ws, f"C{r}", desc, F_BASE, border=False)
    r += 1
ws["B11"].font = F_INPUT
ws["B11"].fill = FILL_INPUT
ws["B12"].font = F_BASE
ws["B13"].font = F_LINK

# ============================================================ 2) الديون
ws = wb.create_sheet("الديون")
ws.sheet_view.rightToLeft = True
ws.sheet_properties.tabColor = TEAL_DARK
title_row(ws, "سجل الديون")
style(ws, "A3", "سجّل كل ديونك هنا — الخلايا الصفراء فقط. الأعمدة الأخرى تُحسب تلقائياً.", F_NOTE, border=False)

headers = ["اسم الدين", "المبلغ المتبقي", "الفائدة السنوية %", "القسط الأدنى",
           "الفائدة الشهرية", "ترتيب كرة الثلج", "ترتيب الانهيار الجليدي"]
for i, h in enumerate(headers, 1):
    style(ws, f"{get_column_letter(i)}4", h, F_HEAD, FILL_HEAD, align=CENTER)
widths = [26, 16, 16, 16, 16, 15, 19]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

FIRST, LAST = 5, 4 + N_DEBTS
examples = [
    ("بطاقة ائتمانية (مثال)", 15000, 0.24, 500),
    ("قرض شخصي (مثال)", 80000, 0.06, 1800),
    ("دين عائلي (مثال)", 20000, 0.0, 400),
]
for r in range(FIRST, LAST + 1):
    ex = examples[r - FIRST] if r - FIRST < len(examples) else None
    style(ws, f"A{r}", ex[0] if ex else None, F_INPUT, FILL_INPUT)
    style(ws, f"B{r}", ex[1] if ex else None, F_INPUT, FILL_INPUT, CUR)
    style(ws, f"C{r}", ex[2] if ex else None, F_INPUT, FILL_INPUT, PCT)
    style(ws, f"D{r}", ex[3] if ex else None, F_INPUT, FILL_INPUT, CUR)
    style(ws, f"E{r}", f'=IF($B{r}="","",$B{r}*$C{r}/12)', F_BASE, None, CUR)
    style(ws, f"F{r}", f'=IF($B{r}="","",COUNTIF($B${FIRST}:$B${LAST},"<"&$B{r})+1)', F_BASE, align=CENTER)
    style(ws, f"G{r}", f'=IF($B{r}="","",COUNTIF($C${FIRST}:$C${LAST},">"&$C{r})+1)', F_BASE, align=CENTER)

T = LAST + 1
style(ws, f"A{T}", "الإجمالي", F_BOLD, FILL_SECTION)
style(ws, f"B{T}", f"=SUM(B{FIRST}:B{LAST})", F_BOLD, FILL_SECTION, CUR)
style(ws, f"C{T}", "", F_BOLD, FILL_SECTION)
style(ws, f"D{T}", f"=SUM(D{FIRST}:D{LAST})", F_BOLD, FILL_SECTION, CUR)
style(ws, f"E{T}", f"=SUM(E{FIRST}:E{LAST})", F_BOLD, FILL_SECTION, CUR)
style(ws, f"F{T}", "", F_BOLD, FILL_SECTION)
style(ws, f"G{T}", "", F_BOLD, FILL_SECTION)
style(ws, f"A{T+2}",
      "كرة الثلج: ابدأ بالترتيب 1 في عمود «كرة الثلج» (الأصغر مبلغاً) لانتصارات سريعة · "
      "الانهيار الجليدي: ابدأ بالترتيب 1 في عموده (الفائدة الأعلى) لأكبر توفير. الأمثلة الثلاثة أعلاه استرشادية — استبدلها بديونك.",
      F_NOTE, border=False)
ws.freeze_panes = "A5"
DEBT_TOTAL_ROW = T

# ============================================================ 3) خطة السداد
ws = wb.create_sheet("خطة السداد")
ws.sheet_view.rightToLeft = True
ws.sheet_properties.tabColor = TEAL_DARK
title_row(ws, "خطة السداد", "J")
for col, w in zip("ABCDEFGHIJ", [22, 16, 14, 14, 16, 3, 16, 14, 14, 16]):
    ws.column_dimensions[col].width = w

style(ws, "A2", "إجمالي الديون", F_BOLD)
style(ws, "B2", f"='الديون'!B{DEBT_TOTAL_ROW}", F_LINK_BOLD, None, CUR)
style(ws, "A3", "مجموع الأقساط الدنيا", F_BOLD)
style(ws, "B3", f"='الديون'!D{DEBT_TOTAL_ROW}", F_LINK, None, CUR)
style(ws, "A4", "الدفعة الشهرية الإضافية ⬅ عدّلها", F_BOLD)
style(ws, "B4", 1000, F_INPUT, FILL_INPUT, CUR)
style(ws, "A5", "الدفعة الشهرية الكلية", F_BOLD)
style(ws, "B5", "=B3+B4", F_BOLD, FILL_SECTION, CUR)
style(ws, "A6", "متوسط الفائدة السنوية المرجّح", F_BOLD)
style(ws, "B6", f"=IFERROR(SUMPRODUCT('الديون'!B{FIRST}:B{LAST},'الديون'!C{FIRST}:C{LAST})/B2,0)", F_BASE, None, PCT)
style(ws, "A7", "الفائدة الشهرية", F_BOLD)
style(ws, "B7", "=B6/12", F_BASE, None, "0.000%")
style(ws, "A8", "أشهر حتى التحرر", F_BOLD)
style(ws, "B8", f'=IF(B2=0,0,IFERROR(MATCH(0,E{SCHED_START}:E{SCHED_END},0),"أكثر من {N_MONTHS}"))',
      F_BOLD, FILL_SECTION, align=CENTER)
style(ws, "A9", "تاريخ التحرر المتوقع", F_BOLD)
style(ws, "B9", '=IF(ISNUMBER(B8),EDATE(TODAY(),B8),"—")', F_BOLD, FILL_SECTION, DATE, align=CENTER)
style(ws, "A10", "إجمالي الفوائد المتوقعة", F_BOLD)
style(ws, "B10", f"=SUM(C{SCHED_START}:C{SCHED_END})", F_BASE, None, CUR)
style(ws, "C2", f'=IF(AND(B2>0,B5<=B2*B7),"⚠ الدفعة لا تغطي الفوائد — ارفع الدفعة الإضافية","")',
      F_WARN, border=False)
ws.merge_cells("C2:E2")
style(ws, "D4", "النموذج تقريبي: يعامل الديون ككتلة واحدة بمتوسط فائدة مرجّح ثابت. "
      "الجدول الأيسر يقارن بسيناريو «الحد الأدنى فقط» (بدون الدفعة الإضافية).", F_NOTE, border=False)
ws.merge_cells("D4:J4")

style(ws, "A11", "الشهر", F_HEAD, FILL_HEAD, align=CENTER)
for col, h in zip("BCDE", ["الرصيد أول الشهر", "الفائدة", "الدفعة", "الرصيد آخر الشهر"]):
    style(ws, f"{col}11", h, F_HEAD, FILL_HEAD, align=CENTER)
style(ws, "G10", "سيناريو الحد الأدنى فقط (للمقارنة)", F_SECTION, FILL_SECTION, align=CENTER)
ws.merge_cells("G10:J10")
for col, h in zip("GHIJ", ["الرصيد أول الشهر", "الفائدة", "الدفعة", "الرصيد آخر الشهر"]):
    style(ws, f"{col}11", h, F_HEAD, FILL_HEAD, align=CENTER)

for i in range(N_MONTHS):
    r = SCHED_START + i
    alt = FILL_ALT if i % 2 else None
    style(ws, f"A{r}", i + 1, F_BASE, alt, align=CENTER)
    style(ws, f"B{r}", "=B2" if i == 0 else f"=E{r-1}", F_BASE, alt, CUR)
    style(ws, f"C{r}", f"=B{r}*$B$7", F_BASE, alt, CUR)
    style(ws, f"D{r}", f"=MIN(B{r}+C{r},$B$5)", F_BASE, alt, CUR)
    style(ws, f"E{r}", f"=MAX(0,B{r}+C{r}-D{r})", F_BASE, alt, CUR)
    style(ws, f"G{r}", "=B2" if i == 0 else f"=J{r-1}", F_BASE, alt, CUR)
    style(ws, f"H{r}", f"=G{r}*$B$7", F_BASE, alt, CUR)
    style(ws, f"I{r}", f"=MIN(G{r}+H{r},$B$3)", F_BASE, alt, CUR)
    style(ws, f"J{r}", f"=MAX(0,G{r}+H{r}-I{r})", F_BASE, alt, CUR)
ws.freeze_panes = "A12"

# ============================================================ 4) الميزانية الشهرية
ws = wb.create_sheet("الميزانية الشهرية")
ws.sheet_view.rightToLeft = True
ws.sheet_properties.tabColor = TEAL_DARK
title_row(ws, "الميزانية الشهرية", "D")
for col, w in zip("ABCD", [30, 16, 16, 16]):
    ws.column_dimensions[col].width = w

style(ws, "A3", "البند", F_HEAD, FILL_HEAD, align=CENTER)
style(ws, "B3", "المخطط", F_HEAD, FILL_HEAD, align=CENTER)
style(ws, "C3", "الفعلي", F_HEAD, FILL_HEAD, align=CENTER)
style(ws, "D3", "الفرق", F_HEAD, FILL_HEAD, align=CENTER)

def section(ws, r, name):
    style(ws, f"A{r}", name, F_SECTION, FILL_SECTION)
    for col in "BCD":
        style(ws, f"{col}{r}", "", F_SECTION, FILL_SECTION)

def budget_row(ws, r, name, planned=None, link=None):
    style(ws, f"A{r}", name, F_BASE)
    if link:
        style(ws, f"B{r}", link, F_LINK, None, CUR)
    else:
        style(ws, f"B{r}", planned, F_INPUT, FILL_INPUT, CUR)
    style(ws, f"C{r}", None, F_INPUT, FILL_INPUT, CUR)
    style(ws, f"D{r}", f'=IF(C{r}="","",B{r}-C{r})', F_BASE, None, CUR)

r = 4
section(ws, r, "الدخل")
income_rows = []
for name, val in [("الراتب الأساسي (مثال)", 18000), ("دخل إضافي", None)]:
    r += 1
    budget_row(ws, r, name, val)
    income_rows.append(r)
r += 1
INCOME_TOTAL = r
style(ws, f"A{r}", "إجمالي الدخل", F_BOLD, FILL_SECTION)
style(ws, f"B{r}", f"=SUM(B{income_rows[0]}:B{income_rows[-1]})", F_BOLD, FILL_SECTION, CUR)
style(ws, f"C{r}", f"=SUM(C{income_rows[0]}:C{income_rows[-1]})", F_BOLD, FILL_SECTION, CUR)
style(ws, f"D{r}", "", F_BOLD, FILL_SECTION)

groups = [
    ("الأساسيات", [("السكن / الإيجار", 4000), ("الفواتير والخدمات", 800), ("الغذاء", 2000),
                    ("المواصلات والوقود", 700), ("الاتصالات", 200), ("أخرى أساسية", 300)]),
    ("الكماليات", [("مطاعم وقهوة", 600), ("ترفيه", 300), ("تسوق", 400), ("اشتراكات", 100)]),
]
group_totals = {}
for gname, items in groups:
    r += 2
    section(ws, r, gname)
    first = r + 1
    for name, val in items:
        r += 1
        budget_row(ws, r, name, val)
    r += 1
    group_totals[gname] = r
    style(ws, f"A{r}", f"إجمالي {gname}", F_BOLD, FILL_SECTION)
    style(ws, f"B{r}", f"=SUM(B{first}:B{r-1})", F_BOLD, FILL_SECTION, CUR)
    style(ws, f"C{r}", f"=SUM(C{first}:C{r-1})", F_BOLD, FILL_SECTION, CUR)
    style(ws, f"D{r}", "", F_BOLD, FILL_SECTION)

r += 2
section(ws, r, "الالتزامات المالية")
first = r + 1
r += 1
budget_row(ws, r, "الأقساط الدنيا للديون", link="='خطة السداد'!B3")
r += 1
budget_row(ws, r, "الدفعة الإضافية (هجوم على الدين)", link="='خطة السداد'!B4")
r += 1
budget_row(ws, r, "ادخار للطوارئ", 500)
r += 1
FIN_TOTAL = r
style(ws, f"A{r}", "إجمالي الالتزامات", F_BOLD, FILL_SECTION)
style(ws, f"B{r}", f"=SUM(B{first}:B{r-1})", F_BOLD, FILL_SECTION, CUR)
style(ws, f"C{r}", f"=SUM(C{first}:C{r-1})", F_BOLD, FILL_SECTION, CUR)
style(ws, f"D{r}", "", F_BOLD, FILL_SECTION)

r += 2
style(ws, f"A{r}", "المتبقي بعد كل المصروفات", F_BOLD, FILL_HEAD)
ws[f"A{r}"].font = F_HEAD
ess, lux = group_totals["الأساسيات"], group_totals["الكماليات"]
style(ws, f"B{r}", f"=B{INCOME_TOTAL}-B{ess}-B{lux}-B{FIN_TOTAL}", F_HEAD, FILL_HEAD, CUR)
style(ws, f"C{r}", f'=IF(C{INCOME_TOTAL}=0,"",C{INCOME_TOTAL}-C{ess}-C{lux}-C{FIN_TOTAL})', F_HEAD, FILL_HEAD, CUR)
style(ws, f"D{r}", "", F_HEAD, FILL_HEAD)
r += 2
style(ws, f"A{r}", "نسبة الأساسيات من الدخل", F_BASE, border=False)
style(ws, f"B{r}", f"=IFERROR(B{ess}/B{INCOME_TOTAL},0)", F_BASE, None, PCT, border=False)
style(ws, f"C{r}", "الهدف الإرشادي ≈ 50%", F_NOTE, border=False)
r += 1
style(ws, f"A{r}", "نسبة الكماليات من الدخل", F_BASE, border=False)
style(ws, f"B{r}", f"=IFERROR(B{lux}/B{INCOME_TOTAL},0)", F_BASE, None, PCT, border=False)
style(ws, f"C{r}", "أثناء سداد الديون حاول ألا تتجاوز 15–20%", F_NOTE, border=False)
r += 1
style(ws, f"A{r}", "نسبة الالتزامات المالية من الدخل", F_BASE, border=False)
style(ws, f"B{r}", f"=IFERROR(B{FIN_TOTAL}/B{INCOME_TOTAL},0)", F_BASE, None, PCT, border=False)
style(ws, f"C{r}", "كل ريال إضافي هنا يقرّب تاريخ حريتك", F_NOTE, border=False)
r += 2
style(ws, f"A{r}", "الأرقام المخططة أعلاه أمثلة استرشادية لدخل 18,000 ر.س — عدّلها لتناسب وضعك.", F_NOTE, border=False)
ws.freeze_panes = "A4"

# ============================================================ 5) لوحة المتابعة
ws = wb.create_sheet("لوحة المتابعة")
ws.sheet_view.rightToLeft = True
ws.sheet_properties.tabColor = TEAL
title_row(ws, "لوحة المتابعة", "F")
for col, w in zip("ABCDEF", [22, 18, 18, 18, 18, 14]):
    ws.column_dimensions[col].width = w

kpis = [
    ("إجمالي الديون", f"='الديون'!B{DEBT_TOTAL_ROW}", CUR),
    ("الدفعة الشهرية الكلية", "='خطة السداد'!B5", CUR),
    ("أشهر حتى التحرر", "='خطة السداد'!B8", None),
    ("تاريخ التحرر المتوقع", "='خطة السداد'!B9", DATE),
    ("إجمالي الفوائد المتوقعة", "='خطة السداد'!B10", CUR),
    ("الفوائد الموفَّرة مقارنة بالحد الأدنى",
     f"=MAX(0,SUM('خطة السداد'!H{SCHED_START}:H{SCHED_END})-SUM('خطة السداد'!C{SCHED_START}:C{SCHED_END}))", CUR),
]
r = 3
for i, (label, formula, fmt) in enumerate(kpis):
    col = "AB CD EF".split()[i % 3]
    if i % 3 == 0 and i > 0:
        r += 3
    c1, c2 = col[0], col[1]
    style(ws, f"{c1}{r}", label, F_HEAD, FILL_HEAD, align=CENTER)
    style(ws, f"{c1}{r+1}", formula, F_LINK_BOLD, FILL_SECTION, fmt, align=CENTER)
    ws.merge_cells(f"{c1}{r}:{c2}{r}")
    ws.merge_cells(f"{c1}{r+1}:{c2}{r+1}")
    ws.row_dimensions[r + 1].height = 26

r += 4
style(ws, f"A{r}", "سجل السداد الفعلي — عبّئ عموديّ الشهر والمبلغ كل شهر", F_SECTION, border=False)
r += 1
LOG_HEAD = r
for col, h in zip("ABCDE", ["الشهر (مثال: 2026/08)", "المبلغ المسدَّد", "التراكمي", "المتبقي من الدين", "نسبة الإنجاز"]):
    style(ws, f"{col}{r}", h, F_HEAD, FILL_HEAD, align=CENTER)
LOG_FIRST = r + 1
LOG_LAST = LOG_FIRST + 23
for i in range(24):
    rr = LOG_FIRST + i
    alt = FILL_ALT if i % 2 else None
    style(ws, f"A{rr}", None, F_INPUT, FILL_INPUT, align=CENTER)
    style(ws, f"B{rr}", None, F_INPUT, FILL_INPUT, CUR)
    style(ws, f"C{rr}", f'=IF(B{rr}="","",SUM($B${LOG_FIRST}:B{rr}))', F_BASE, alt, CUR)
    style(ws, f"D{rr}", f'=IF(B{rr}="","",MAX(0,\'الديون\'!$B${DEBT_TOTAL_ROW}-C{rr}))', F_BASE, alt, CUR)
    style(ws, f"E{rr}", f'=IF(B{rr}="","",MIN(1,C{rr}/\'الديون\'!$B${DEBT_TOTAL_ROW}))', F_BOLD, alt, PCT, align=CENTER)
ws.conditional_formatting.add(
    f"E{LOG_FIRST}:E{LOG_LAST}",
    DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                color=TEAL, showValue=True, minLength=None, maxLength=None),
)
style(ws, f"A{LOG_LAST+2}",
      "ملاحظة: «المتبقي من الدين» يفترض أن إجمالي ورقة الديون ثابت كما أدخلته أول مرة — حدّث ورقة الديون كل عدة أشهر لدقة أعلى.",
      F_NOTE, border=False)
ws.freeze_panes = f"A{LOG_FIRST}"

wb.save(OUT)
print(f"saved: {OUT}")
